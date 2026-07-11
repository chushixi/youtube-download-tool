"""把比較結果寫成 Excel。"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .wears import WEAR_BY_KEY

# 欄位定義：(標題, 取值 key, 寬度, 數字格式)
COLUMNS = [
    ("皮膚",           "skin",              22, None),
    ("磨損",           "wear_zh",           10, None),
    ("market_hash_name", "market_hash_name", 34, None),
    ("Steam最低賣價",  "steam_lowest_sell", 14, "#,##0.00"),
    ("Steam最高求購",  "steam_highest_buy", 14, "#,##0.00"),
    ("Steam交易量(24h)", "steam_volume",    14, "#,##0"),
    ("BUFF最低賣價",   "buff_lowest_sell",  14, "#,##0.00"),
    ("BUFF最高求購",   "buff_highest_buy",  14, "#,##0.00"),
    ("BUFF在售量",     "buff_sell_num",     12, "#,##0"),
    ("賣價差(Steam-BUFF)", "sell_diff",     16, "#,##0.00"),
    ("賣價比(Steam/BUFF)", "sell_ratio",    16, "0.00"),
    ("抓取時間",       "fetched_at",        20, None),
]

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_THIN = Side(style="thin", color="D9D9D9")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_ALT_FILL = PatternFill("solid", fgColor="F2F6FB")


def write_workbook(rows, path, currency_label="¥ (人民幣)"):
    """rows: pipeline.build_rows() 產出的 dict 清單。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "AK-47 價格比較"

    # 說明列
    ws.cell(row=1, column=1,
            value=f"CS2 AK-47 皮膚 Steam vs BUFF 價格比較　幣別：{currency_label}")
    ws.cell(row=1, column=1).font = Font(bold=True, size=12)
    ws.merge_cells(start_row=1, start_column=1, end_row=1,
                   end_column=len(COLUMNS))

    header_row = 2
    for c, (title, _key, width, _fmt) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=c, value=title)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)
        cell.border = _BORDER
        ws.column_dimensions[get_column_letter(c)].width = width

    for i, row in enumerate(rows):
        r = header_row + 1 + i
        for c, (_title, key, _width, fmt) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=r, column=c, value=row.get(key))
            cell.border = _BORDER
            if fmt:
                cell.number_format = fmt
            if i % 2 == 1:
                cell.fill = _ALT_FILL

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    ws.auto_filter.ref = (
        f"A{header_row}:{get_column_letter(len(COLUMNS))}{header_row}")

    wb.save(path)
    return path


def wear_zh(wear_key):
    w = WEAR_BY_KEY.get(wear_key)
    return w["zh_tw"] if w else wear_key
